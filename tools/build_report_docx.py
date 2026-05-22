from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor


OUTPUT = "Relatorio_Trabalho09_Tratamento_de_Restricoes.docx"


def set_run(run, bold=False, size=11, color=None):
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(document, text, level=1):
    paragraph = document.add_heading(level=level)
    run = paragraph.add_run(text)
    set_run(run, bold=True, size=16 if level == 1 else 13, color=(31, 78, 94))
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_body(document, text):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run(run, size=11)
    paragraph.paragraph_format.line_spacing = 1.12
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_code_block(document, lines):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Courier New"
    run.font.size = Pt(9.5)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    table.style = "Table Grid"
    return table


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].text = header
        header_cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in header_cells[index].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run(run, bold=True, size=10, color=(31, 78, 94))

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run(run, size=10)

    document.add_paragraph()
    return table


def add_print_placeholder(document, title, filename):
    add_heading(document, title, level=2)
    add_body(document, "Inserir aqui o print correspondente depois que o aplicativo estiver rodando.")
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"ESPACO PARA PRINT\n{filename}")
    set_run(run, bold=True, size=12, color=(69, 100, 107))
    for _ in range(5):
        paragraph.add_run("\n")
    document.add_paragraph()


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Relatorio - Trabalho 09\nTratamento de Restricoes")
    set_run(run, bold=True, size=20, color=(31, 78, 94))

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Mateus Flausino Conceicao | 12612EEL011")
    set_run(run, size=11, color=(69, 100, 107))

    add_heading(document, "1. Objetivo")
    add_body(
        document,
        "Este trabalho desenvolve uma aplicacao interativa para demonstrar o uso de Algoritmo Genetico em um problema de otimizacao restrita. O desafio indicado no material de apoio consiste em realizar 30 execucoes do AG e obter a media e o desvio padrao dos resultados.",
    )

    add_heading(document, "2. Problema de otimizacao")
    add_body(document, "O problema implementado busca minimizar:")
    add_code_block(
        document,
        [
            "f(x) = (x1 - 1)^2 + (x2 - 2)^2 + 1",
            "sujeito a:",
            "x1 + x2 <= 4",
            "x1 >= 2",
            "x2 >= 1",
            "0 <= x1, x2 <= 5",
        ],
    )
    add_body(document, "O otimo factivel esperado esta em x = (2, 2), com valor f(x) = 2.")

    add_heading(document, "3. Metodo")
    add_body(
        document,
        "Foram implementados dois modos de tratamento de restricoes: factivel primeiro, baseado no criterio de Deb, e penalizacao estatica. O AG utiliza representacao real, selecao por roleta ponderada por ranking, crossover aritmetico e mutacao gaussiana.",
    )

    add_heading(document, "4. Parametros utilizados")
    add_table(
        document,
        ["Parametro", "Valor"],
        [
            ["Tamanho da populacao", "60"],
            ["Taxa de crossover", "0.78"],
            ["Taxa de mutacao", "0.16"],
            ["Escala da mutacao", "0.35"],
            ["Numero de geracoes", "180"],
            ["Numero de execucoes", "30"],
        ],
    )

    add_heading(document, "5. Prints do aplicativo")
    add_print_placeholder(document, "Print 1 - Tela inicial do app", "prints/print-01-tela-inicial.png")
    add_print_placeholder(document, "Print 2 - Evolucao da populacao", "prints/print-02-evolucao.png")
    add_print_placeholder(document, "Print 3 - Resultado das 30 execucoes", "prints/print-03-trinta-execucoes.png")

    add_heading(document, "6. Resultados")
    add_table(
        document,
        ["Indicador", "Valor observado"],
        [
            ["Media dos melhores valores de f(x)", ""],
            ["Desvio padrao dos melhores valores de f(x)", ""],
            ["Taxa de execucoes factiveis", ""],
            ["Melhor valor obtido", ""],
            ["Pior valor obtido", ""],
        ],
    )

    add_heading(document, "7. Discussao")
    add_body(
        document,
        "O criterio factivel primeiro tende a conduzir a populacao para a regiao valida antes de intensificar a busca no valor da funcao objetivo. A penalizacao estatica tambem pode funcionar, mas depende diretamente do peso de penalidade escolhido.",
    )

    add_heading(document, "8. Conclusao")
    add_body(
        document,
        "A aplicacao permite visualizar o efeito das restricoes no processo evolutivo e cumpre o desafio de executar o AG 30 vezes, calculando media e desvio padrao dos resultados.",
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    build()
